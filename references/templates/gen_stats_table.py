#!/usr/bin/env python3
"""生成多站点渗透测试报告·漏洞统计表格 docx（独立文件，供用户粘贴进 DLP 加密的主报告）。
用法：PYTHONPATH=/home/kali/.local/lib/python3.13/site-packages python3 gen_stats_table.py
数据源：把 rows 数组换成实际统计即可。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run('各站点漏洞数量统计')
r.bold = True
r.font.size = Pt(14)

# (站点, 高危, 中危, 低危, 合计) —— 换成实际数据
rows = [
    ('站点', '高危', '中危', '低危', '合计'),
    ('一、<资产A>',        1, 3, 2, 6),
    ('二、<资产B>',        0, 3, 1, 4),
    ('合计',               1, 6, 3, 10),
]

table = doc.add_table(rows=len(rows), cols=5)
table.style = 'Table Grid'
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = table.cell(i, j)
        cell.text = str(val)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            run.font.size = Pt(10.5)
            run.font.name = 'Calibri'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            if i == 0 or row[0] == '合计':
                run.bold = True

note = doc.add_paragraph()
nr = note.add_run('注：等级按各漏洞实际影响评定——任意密码重置/管理员接管/批量PII/任意文件读取密码/匿名写存储计高危；信息泄露、越权读取、爆破前提类计中危；用户枚举、点击劫持、纯结构泄露计低危。')
nr.font.size = Pt(9)
nr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.save('各站点漏洞数量统计表.docx')
print('written: 各站点漏洞数量统计表.docx')
